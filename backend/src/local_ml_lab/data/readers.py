import csv
import re
import unicodedata
from pathlib import Path

import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook

from local_ml_lab.settings import settings

SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".parquet"}
CSV_ENCODINGS = {"utf-8", "utf-8-sig", "latin-1"}
CSV_DELIMITERS = {",", ";", "\t", "|"}


def inspect_file(path: Path) -> dict:
    """Inspect metadata without evaluating formulas or loading a full workbook."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("UNSUPPORTED_FILE_TYPE")
    result = {"extension": suffix, "size_bytes": path.stat().st_size}
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            result["sheet_names"] = workbook.sheetnames
            result["sheets"] = [
                {
                    "name": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                }
                for sheet in workbook.worksheets
            ]
        finally:
            workbook.close()
    elif suffix == ".parquet":
        metadata = pq.ParquetFile(path).metadata
        result.update({"row_count": metadata.num_rows, "column_count": metadata.num_columns})
    else:
        encoding = _detect_encoding(path)
        sample = path.read_bytes()[:65536].decode(encoding)
        result.update(
            {
                "encoding": encoding,
                "delimiter": _detect_delimiter(sample),
                "has_header": True,
            }
        )
    return result


def read_frame(path: Path, options: dict | None = None) -> pd.DataFrame:
    path = Path(path)
    options = options or {}
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("UNSUPPORTED_FILE_TYPE")
    header_row = _header_row(options.get("header_row", 0))
    if suffix == ".csv":
        encoding = options.get("encoding") or _detect_encoding(path)
        if encoding not in CSV_ENCODINGS:
            raise ValueError("UNSUPPORTED_CSV_ENCODING")
        sample = path.read_bytes()[:65536].decode(encoding)
        delimiter = options.get("delimiter") or _detect_delimiter(sample)
        if delimiter not in CSV_DELIMITERS:
            raise ValueError("UNSUPPORTED_CSV_DELIMITER")
        frame = pd.read_csv(
            path,
            encoding=encoding,
            sep=delimiter,
            header=header_row,
            nrows=settings.max_rows + 1,
            low_memory=False,
        )
    elif suffix == ".xlsx":
        metadata = inspect_file(path)
        sheets = metadata.get("sheet_names", [])
        requested = options.get("sheet_name")
        if requested is None and len(sheets) > 1:
            raise ValueError("XLSX_SHEET_REQUIRED")
        sheet_name = requested if requested is not None else sheets[0] if sheets else 0
        if isinstance(sheet_name, str) and sheet_name not in sheets:
            raise ValueError("XLSX_SHEET_NOT_FOUND")
        frame = pd.read_excel(
            path,
            sheet_name=sheet_name,
            header=header_row,
            nrows=settings.max_rows + 1,
            engine="openpyxl",
        )
    else:
        if options:
            unsupported = set(options) - {"columns"}
            if unsupported:
                raise ValueError("UNSUPPORTED_PARQUET_OPTIONS")
        columns = options.get("columns")
        frame = pd.read_parquet(path, columns=columns)
        if len(frame) > settings.max_rows:
            frame = frame.iloc[: settings.max_rows + 1]
    _validate_limits(frame)
    if frame.columns.empty or all(str(column).startswith("Unnamed:") for column in frame.columns):
        raise ValueError("HEADER_ROW_NOT_FOUND")
    return frame


def normalize_columns(frame: pd.DataFrame) -> tuple[pd.DataFrame, list[dict]]:
    counts: dict[str, int] = {}
    mapping = []
    for index, raw in enumerate(frame.columns, 1):
        base = _display_name(raw)
        counts[base] = counts.get(base, 0) + 1
        display = base if counts[base] == 1 else f"{base} ({counts[base]})"
        mapping.append(
            {
                "column_id": f"c{index:04d}",
                "display_name": display,
                "source_name": str(raw),
                "position": index - 1,
            }
        )
    normalized = frame.copy()
    normalized.columns = [row["column_id"] for row in mapping]
    return normalized, mapping


def profile_frame(
    frame: pd.DataFrame,
    mapping: list[dict],
    type_overrides: dict[str, str] | None = None,
    roles: dict[str, str] | None = None,
) -> dict:
    type_overrides = type_overrides or {}
    roles = roles or {}
    sampled = len(frame) > settings.max_profile_sample_rows
    sample = frame.head(settings.max_profile_sample_rows) if sampled else frame
    columns = []
    for meta in mapping:
        series = sample[meta["column_id"]]
        non_null = series.dropna()
        distinct = int(non_null.nunique(dropna=True))
        semantic = type_overrides.get(meta["column_id"]) or _semantic_type(non_null)
        possible_id = _possible_id(non_null, semantic, meta["display_name"])
        issues = []
        if series.isna().any():
            issues.append("MISSING_VALUES")
        if distinct <= 1:
            issues.append("CONSTANT_OR_EMPTY")
        if possible_id:
            issues.append("POSSIBLE_IDENTIFIER")
        columns.append(
            {
                **meta,
                "inferred_semantic_type": semantic,
                "null_count": int(series.isna().sum()),
                "distinct_count": distinct,
                "possible_id": possible_id,
                "configured_role": roles.get(meta["column_id"], "variable"),
                "examples": [_safe_example(value) for value in non_null.head(4)],
                "quality_issue_codes": issues,
            }
        )
    return {
        "row_count": len(frame),
        "column_count": len(frame.columns),
        "duplicate_count": int(frame.duplicated().sum()),
        "sampled": sampled,
        "profile_sample_count": len(sample),
        "columns": columns,
    }


def _validate_limits(frame: pd.DataFrame) -> None:
    if len(frame) > settings.max_rows:
        raise ValueError("MAX_ROWS_EXCEEDED")
    if len(frame.columns) > settings.max_columns:
        raise ValueError("MAX_COLUMNS_EXCEEDED")
    if len(frame) * len(frame.columns) > settings.max_cells:
        raise ValueError("MAX_CELLS_EXCEEDED")


def _detect_encoding(path: Path) -> str:
    sample = path.read_bytes()[:65536]
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            sample.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return "latin-1"


def _detect_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        raise ValueError("CSV_DELIMITER_NOT_DETECTED") from None


def _header_row(value) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or not 0 <= value <= 99:
        raise ValueError("INVALID_HEADER_ROW")
    return value


def _display_name(value) -> str:
    text = re.sub(r"\s+", " ", str(value).strip())
    return text[:200] or "Sin nombre"


def _semantic_type(series: pd.Series) -> str:
    if pd.api.types.is_bool_dtype(series):
        return "categorical"
    if pd.api.types.is_datetime64_any_dtype(series):
        return "datetime"
    if pd.api.types.is_numeric_dtype(series):
        return "numeric"
    if pd.api.types.is_object_dtype(series) or pd.api.types.is_string_dtype(series):
        text = series.astype(str).str.strip()
        date_like = text.str.match(r"^\d{4}[-/]\d{1,2}([-/]\d{1,2})?$") | text.str.match(
            r"^\d{1,2}[-/]\d{1,2}[-/]\d{2,4}$"
        )
        if len(text) and date_like.mean() >= 0.9:
            return "datetime"
    return "categorical"


def _possible_id(series: pd.Series, semantic: str, display_name: str = "") -> bool:
    if len(series) < 10 or semantic == "datetime":
        return False
    unique_ratio = series.nunique(dropna=True) / max(1, len(series))
    ascii_name = unicodedata.normalize("NFKD", display_name).encode("ascii", "ignore").decode()
    normalized_name = re.sub(r"[^a-z0-9]+", "_", ascii_name.lower()).strip("_")
    id_name = bool(
        re.search(r"(^id$|_id$|^id_|uuid|codigo|code$|numero$|operacion)", normalized_name)
    )
    if semantic == "numeric":
        return id_name and unique_ratio >= 0.98
    average_length = series.astype(str).str.len().mean() if len(series) else 0
    return unique_ratio >= 0.98 and (id_name or average_length >= 8)


def _safe_example(value):
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, str):
        return value.replace("\n", " ").replace("\r", " ")[:120]
    return value
