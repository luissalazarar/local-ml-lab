import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from local_ml_lab.data.preparation import (
    PreparationError,
    apply_recipe,
    canonical_hash,
    canonical_recipe,
    export_prepared_excel,
    frame_hash,
)
from local_ml_lab.data.readers import normalize_columns, profile_frame
from local_ml_lab.domain.contracts import PreparationRecipe


def recipe(column, **options):
    return PreparationRecipe.model_validate({"schema_version": "1.0", "columns": {column: options}})


@pytest.mark.parametrize(
    ("date_format", "value", "expected"),
    [
        ("DMY", "15/03/2026", "2026-03-15"),
        ("MDY", "03/15/2026", "2026-03-15"),
        ("YMD", "2026-03-15", "2026-03-15"),
    ],
)
def test_confirmed_date_formats(date_format, value, expected):
    frame = pd.DataFrame({"c0001": [value]})
    result = apply_recipe(
        frame,
        recipe("c0001", type="date", date_format=date_format),
        source_extension=".xlsx",
    )
    assert result.frame.loc[0, "c0001"].strftime("%Y-%m-%d") == expected


def test_real_excel_datetime_and_serial_are_supported_only_when_confirmed():
    frame = pd.DataFrame({"c0001": [pd.Timestamp("2026-01-02"), 46024]})
    result = apply_recipe(
        frame,
        recipe("c0001", type="date", date_format="YMD", excel_date_serials=True),
        source_extension=".xlsx",
    )
    assert result.frame["c0001"].dt.year.tolist() == [2026, 2026]
    with pytest.raises(PreparationError, match="UNCONVERTIBLE_VALUES"):
        apply_recipe(
            frame,
            recipe("c0001", type="date", date_format="YMD"),
            source_extension=".csv",
        )


def test_ambiguous_date_is_blocked_until_user_confirms():
    frame = pd.DataFrame({"c0001": ["01/02/2026"]})
    with pytest.raises(PreparationError, match="UNCONVERTIBLE_VALUES"):
        apply_recipe(
            frame,
            recipe("c0001", type="date", date_format="UNAMBIGUOUS"),
            source_extension=".xlsx",
        )


def test_invalid_date_is_auditable_quarantine():
    frame = pd.DataFrame({"c0001": ["15/03/2026", "imposible"]})
    result = apply_recipe(
        frame,
        recipe("c0001", type="date", date_format="DMY", invalid="segregate"),
        source_extension=".xlsx",
    )
    assert len(result.frame) == 1
    assert result.quarantined[0]["row_id"] == "row-0000002"
    assert result.quarantined[0]["stage"] == "preparación"


@pytest.mark.parametrize(
    ("values", "options", "expected"),
    [
        (["1.234,50"], {"decimal_separator": ",", "thousands_separator": "."}, 1234.5),
        (["1,234.50"], {"decimal_separator": ".", "thousands_separator": ","}, 1234.5),
        (["15%"], {"decimal_separator": ".", "percent": True}, 0.15),
        (
            ["S/ 1,234.50"],
            {"decimal_separator": ".", "thousands_separator": ",", "currency": "PEN"},
            1234.5,
        ),
    ],
)
def test_confirmed_numeric_representations(values, options, expected):
    result = apply_recipe(
        pd.DataFrame({"c0001": values}),
        recipe("c0001", type="numeric", **options),
        source_extension=".xlsx",
    )
    assert result.frame.loc[0, "c0001"] == expected


def test_mixed_currencies_are_blocked():
    frame = pd.DataFrame({"c0001": ["S/ 10.00", "USD 10.00"]})
    with pytest.raises(PreparationError, match="MIXED_CURRENCIES"):
        apply_recipe(
            frame,
            recipe("c0001", type="numeric", decimal_separator=".", currency="PEN"),
            source_extension=".xlsx",
        )


def test_trim_empty_duplicates_roles_and_original_intact():
    original = pd.DataFrame({"c0001": [" Norte ", "", " Norte "], "c0002": [1, 2, 1]})
    before = original.copy(deep=True)
    configured = PreparationRecipe.model_validate(
        {
            "columns": {
                "c0001": {
                    "type": "categorical",
                    "trim": True,
                    "empty_to_missing": True,
                },
                "c0002": {"role": "identifier"},
            },
            "exact_duplicates": "exclude",
        }
    )
    result = apply_recipe(original, configured, source_extension=".xlsx")
    assert original.equals(before)
    assert result.frame["c0001"].tolist()[0] == "Norte"
    assert pd.isna(result.frame["c0001"].tolist()[1])
    assert result.roles["c0002"] == "identifier"


def test_duplicate_keep_and_exclude_are_explicit():
    frame = pd.DataFrame({"c0001": [1, 1]})
    kept = apply_recipe(frame, PreparationRecipe(), source_extension=".csv")
    excluded = apply_recipe(
        frame, PreparationRecipe(exact_duplicates="exclude"), source_extension=".csv"
    )
    assert len(kept.frame) == 2
    assert len(excluded.frame) == 1
    assert excluded.quarantined[0]["stage"] == "duplicados"


def test_date_filter_and_monthly_sum_mean():
    frame = pd.DataFrame(
        {
            "c0001": pd.to_datetime(["2026-01-01", "2026-01-15", "2026-02-01"]),
            "c0002": [10.0, 20.0, 30.0],
        }
    )
    base = {
        "columns": {"c0001": {"type": "date", "date_format": "YMD"}},
        "filters": [
            {
                "kind": "date_range",
                "column_id": "c0001",
                "start": "2026-01-01",
                "end": "2026-02-28",
            }
        ],
    }
    for operation, first in [("sum", 30.0), ("mean", 15.0)]:
        configured = PreparationRecipe.model_validate(
            {
                **base,
                "monthly_aggregation": {
                    "date_column_id": "c0001",
                    "value_column_id": "c0002",
                    "operation": operation,
                },
            }
        )
        result = apply_recipe(frame, configured, source_extension=".xlsx")
        assert result.frame.loc[0, "c0002"] == first


def test_recipe_and_output_hash_are_deterministic_and_contain_no_learned_statistics():
    frame = pd.DataFrame({"c0001": ["1,5", "2,5"]})
    configured = recipe("c0001", type="numeric", decimal_separator=",")
    first = apply_recipe(frame, configured, source_extension=".xlsx")
    second = apply_recipe(frame, configured, source_extension=".xlsx")
    payload = json.dumps(canonical_recipe(configured), sort_keys=True)
    forbidden = ["mean", "median", "scale", "target_categories", "imputer", "encoder"]
    assert not any(term in payload for term in forbidden)
    assert canonical_hash(canonical_recipe(configured)) == canonical_hash(
        canonical_recipe(configured)
    )
    assert frame_hash(first.frame) == frame_hash(second.frame)


def test_constant_and_identifier_remain_profile_recommendations():
    frame, mapping = normalize_columns(
        pd.DataFrame({"Cliente_ID": [f"C-{index}" for index in range(20)], "Constante": 1})
    )
    columns = profile_frame(frame, mapping)["columns"]
    assert columns[0]["possible_id"] is True
    assert "CONSTANT_OR_EMPTY" in columns[1]["quality_issue_codes"]


def test_prepared_excel_has_required_sheets_and_formula_protection(tmp_path):
    frame = pd.DataFrame({"c0001": ["=1+1", "normal"]})
    mapping = [{"column_id": "c0001", "display_name": "Texto"}]
    profile = profile_frame(frame, [{**mapping[0], "source_name": "Texto", "position": 0}])
    path = tmp_path / "prepared.xlsx"
    export_prepared_excel(
        path,
        frame,
        mapping,
        [],
        [],
        {"original": {"filas": 2}, "prepared": {"filas": 2}},
        {},
    )
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == [
        "01_Datos_Preparados",
        "02_Filas_Apartadas",
        "03_Transformaciones",
        "04_Calidad",
        "05_Diccionario",
    ]
    assert workbook["01_Datos_Preparados"]["A2"].value == "'=1+1"
    assert profile["row_count"] == 2
