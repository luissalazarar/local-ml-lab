from openpyxl import load_workbook
from pypdf import PdfReader

from local_ml_lab.reports import create_excel, create_pdf


def result_fixture():
    return {
        "goal": "estimate_value",
        "problem_type": "regression",
        "primary_metric_id": "mae",
        "engine_version": "test",
        "dataset_summary": {"row_count": 2, "column_count": 2},
        "selection_decision": {
            "model_id": "ridge",
            "reason": "Superó la mejora práctica y la consistencia.",
        },
        "reliability": {"primary_level": "low"},
        "validation_plan": {"evidence_mode": "selection_cv"},
        "evaluation_metrics": [
            {
                "metric_id": "mae",
                "name": "MAE",
                "value": 1.5,
                "n_used": 2,
                "evaluation_role": "selection_oof",
            }
        ],
        "data_quality": {
            "columns": [
                {
                    "column_id": "c0001",
                    "display_name": "ventas",
                    "null_count": 1,
                    "possible_id": False,
                }
            ]
        },
        "resolved_config": {"primary_metric": "mae"},
        "data_preparation": {
            "dataset_version_id": "prepared-test",
            "rows_input": 3,
            "rows_analyzed": 2,
            "rows_quarantined": 1,
            "transformations": [
                {
                    "column_id": "c0001",
                    "transformation": "interpretar como número",
                    "affected_count": 2,
                }
            ],
        },
        "candidates": [
            {"model_id": "ridge", "display_name": "Regresión Ridge", "status": "succeeded"}
        ],
        "predictions": [
            {"actual": 10.0, "predicted": 11.0, "error": 1.0},
            {"actual": 20.0, "predicted": 18.0, "error": -2.0},
        ],
        "drivers": [
            {
                "source_column_id": "c0001",
                "importance_mean": 0.5,
                "evaluation_role": "selection_validation_folds",
            }
        ],
        "limitations": ["Validación usada para selección"],
    }


def test_excel_and_pdf_open_with_real_content(tmp_path):
    excel_path = tmp_path / "report.xlsx"
    pdf_path = tmp_path / "report.pdf"
    create_excel(result_fixture(), excel_path)
    create_pdf(result_fixture(), pdf_path)
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    assert workbook["02_Preparacion"]["A1"].value is not None
    assert "Sin errores" in {
        str(cell.value) for row in workbook["08_Errores"].iter_rows() for cell in row
    }
    assert len(list(workbook["06_Predicciones"].iter_rows(values_only=True))) == 3
    assert workbook["12_Seleccion_Modelo"]["A1"].value is not None
    assert workbook["13_Limites"]["A1"].value is not None
    summary_values = {str(cell.value) for row in workbook["00_Resumen"].iter_rows() for cell in row}
    assert "Desarrollado por Luis Salazar" in summary_values
    assert "Regresión Ridge" in summary_values
    reader = PdfReader(pdf_path)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Métrica principal" in text
    assert "Variables predictivas" in text
    assert "Desarrollado por Luis Salazar" in text
    assert "Regresión Ridge" in text
    assert "Preparación de datos" in text
    assert "Cómo se eligió el modelo" in text
    assert "Cómo se evaluó" in text
    assert "Qué significan los límites" in text
    assert len(pdf_path.read_bytes()) > 2_000
