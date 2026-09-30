from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

root = Path("/data/runs")
xlsx = sorted(root.rglob("*.xlsx"), key=lambda path: path.stat().st_mtime)[-1]
pdf = sorted(root.rglob("*.pdf"), key=lambda path: path.stat().st_mtime)[-1]
workbook = load_workbook(xlsx, read_only=True, data_only=False)
reader = PdfReader(pdf)
text = "".join(page.extract_text() or "" for page in reader.pages)
expected = {"00_Resumen", "05_Metricas", "06_Predicciones", "11_Diccionario"}
assert expected.issubset(workbook.sheetnames)
assert workbook["02_Limpieza"]["A1"].value is not None
assert workbook["08_Errores"]["A1"].value is not None
assert len(list(workbook["06_Predicciones"].iter_rows(values_only=True))) > 1
assert len(reader.pages) >= 2
assert "Limitaciones" in text and "Métrica principal" in text and "Variables predictivas" in text
print({"xlsx_sheets": len(workbook.sheetnames), "pdf_pages": len(reader.pages), "text_verified": True})
